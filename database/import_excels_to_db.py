import os
from datetime import datetime

import bcrypt
import openpyxl
import pymysql

BASE_DIR = r"d:\projects\unified-academic-dashboard-Hari\excels"
DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "12345",
    "database": "campusone",
    "charset": "utf8mb4",
    "autocommit": False,
}

DEFAULT_PASSWORD = "Campus@123"


def norm(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def norm_int(v):
    s = norm(v)
    if s is None:
        return None
    try:
        return int(float(s))
    except Exception:
        return None


def normalize_room_type(v):
    s = (norm(v) or "").upper()
    return "LAB" if "LAB" in s else "CLASS"


def normalize_course_type(v):
    s = (norm(v) or "THEORY").upper()
    return "LAB" if s == "LAB" else "THEORY"


def normalize_dept_code(v):
    s = (norm(v) or "").upper()
    if s in ("BIOTECH", "BIOTECHNOLOGY"):
        return "BIO"
    return s


def get_sheet_rows(file_name):
    path = os.path.join(BASE_DIR, file_name)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rows.append({headers[i]: row_vals[i] if i < len(row_vals) else None for i in range(len(headers))})
    return rows


def ensure_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS faculty_courses (
            faculty_id VARCHAR(255) NOT NULL,
            course_code VARCHAR(255) NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (faculty_id, course_code),
            INDEX idx_faculty_courses_code (course_code),
            CONSTRAINT fk_faculty_courses_faculty
                FOREIGN KEY (faculty_id) REFERENCES faculty(faculty_id)
                ON DELETE CASCADE
        ) ENGINE=InnoDB
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS room_allocations (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            room_number VARCHAR(100) NOT NULL,
            room_type ENUM('CLASS', 'LAB') DEFAULT 'CLASS',
            department VARCHAR(100),
            semester INT,
            section VARCHAR(20),
            course_code VARCHAR(40),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE KEY uniq_room_alloc (room_number, department, semester, section, course_code),
            INDEX idx_room_alloc_course (course_code),
            CONSTRAINT fk_room_alloc_room
                FOREIGN KEY (room_number) REFERENCES rooms(room_number)
                ON DELETE CASCADE
        ) ENGINE=InnoDB
        """
    )


def load_departments(cur):
    cur.execute("SELECT id, department_code, department_name FROM departments")
    by_code = {}
    by_name = {}
    for dep_id, dep_code, dep_name in cur.fetchall():
        if dep_code:
            by_code[str(dep_code).strip().upper()] = dep_id
        if dep_name:
            by_name[str(dep_name).strip().upper()] = dep_id
    return by_code, by_name


def get_or_create_account(cur, *, role, full_name, email, department, semester, academic_year, section, age, mobile, address, blood_group):
    email_norm = norm(email)
    if not email_norm:
        return None, "missing email"
    email_norm = email_norm.lower()

    cur.execute("SELECT id, role FROM accounts WHERE email=%s", (email_norm,))
    row = cur.fetchone()
    if row:
        account_id, existing_role = row
        if existing_role not in (role,):
            return None, f"email {email_norm} belongs to {existing_role}"
        cur.execute(
            """
            UPDATE accounts
            SET full_name=%s,
                department=%s,
                semester=%s,
                academic_year=%s,
                age=%s,
                mobile_number=%s,
                address=%s,
                blood_group=%s
            WHERE id=%s
            """,
            (
                norm(full_name),
                norm(department),
                semester,
                norm(academic_year),
                age,
                norm(mobile),
                norm(address),
                norm(blood_group),
                account_id,
            ),
        )
        return account_id, None

    pwd_hash = bcrypt.hashpw(DEFAULT_PASSWORD.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
    cur.execute(
        """
        INSERT INTO accounts (
            full_name, email, password, role, status, department, semester,
            academic_year, age, mobile_number, address, blood_group, created_at
        ) VALUES (%s,%s,%s,%s,'ACTIVE',%s,%s,%s,%s,%s,%s,%s,NOW())
        """,
        (
            norm(full_name),
            email_norm,
            pwd_hash,
            role,
            norm(department),
            semester,
            norm(academic_year),
            age,
            norm(mobile),
            norm(address),
            norm(blood_group),
        ),
    )
    return cur.lastrowid, None


def import_faculty(cur):
    rows = get_sheet_rows("Faculty.xlsx")
    ok, skip = 0, 0

    for r in rows:
        faculty_id = norm(r.get("FacultyID"))
        if not faculty_id:
            skip += 1
            continue

        account_id, err = get_or_create_account(
            cur,
            role="FACULTY",
            full_name=r.get("Name"),
            email=r.get("Email"),
            department=normalize_dept_code(r.get("Department")),
            semester=None,
            academic_year=None,
            section=None,
            age=norm_int(r.get("Age")),
            mobile=r.get("Mobile"),
            address=r.get("Address"),
            blood_group=r.get("BloodGroup"),
        )

        if err or not account_id:
            skip += 1
            continue

        cur.execute("SELECT account_id FROM faculty WHERE faculty_id=%s", (faculty_id,))
        existing = cur.fetchone()
        if existing and existing[0] != account_id:
            skip += 1
            continue

        cur.execute(
            """
            INSERT INTO faculty (account_id, faculty_id)
            VALUES (%s, %s)
            ON DUPLICATE KEY UPDATE faculty_id=VALUES(faculty_id)
            """,
            (account_id, faculty_id),
        )
        ok += 1

    return ok, skip


def import_students(cur):
    rows = get_sheet_rows("students.xlsx")
    ok, skip = 0, 0

    for r in rows:
        reg_no = norm(r.get("RegNo"))
        if not reg_no:
            skip += 1
            continue

        account_id, err = get_or_create_account(
            cur,
            role="STUDENT",
            full_name=r.get("Name"),
            email=r.get("Email"),
            department=normalize_dept_code(r.get("Department")),
            semester=norm_int(r.get("Semester")),
            academic_year=norm(r.get("AcademicYear")),
            section=norm(r.get("Section")),
            age=norm_int(r.get("Age")),
            mobile=r.get("Mobile"),
            address=r.get("Address"),
            blood_group=r.get("BloodGroup"),
        )

        if err or not account_id:
            skip += 1
            continue

        cur.execute("SELECT account_id FROM students WHERE registration_number=%s", (reg_no,))
        existing = cur.fetchone()
        if existing and existing[0] != account_id:
            skip += 1
            continue

        cur.execute(
            """
            INSERT INTO students (account_id, registration_number, section)
            VALUES (%s, %s, %s)
            ON DUPLICATE KEY UPDATE registration_number=VALUES(registration_number), section=VALUES(section)
            """,
            (account_id, reg_no, norm(r.get("Section"))),
        )
        ok += 1

    return ok, skip


def import_courses(cur):
    rows = get_sheet_rows("course.xlsx")
    dep_by_code, dep_by_name = load_departments(cur)
    ok, skip = 0, 0

    for r in rows:
        code = norm(r.get("CourseCode"))
        name = norm(r.get("CourseName"))
        semester = norm_int(r.get("Semester"))
        credits = norm_int(r.get("Credits"))
        year = norm(r.get("Year"))
        dept_raw = normalize_dept_code(r.get("Department"))

        if not code or not name or semester is None or credits is None or not dept_raw:
            skip += 1
            continue

        dep_id = dep_by_code.get(dept_raw.upper()) or dep_by_name.get(dept_raw.upper())
        if not dep_id:
            skip += 1
            continue

        if year and len(year) == 4 and year.isdigit():
            academic_year = f"{year} - {int(year) + 4}"
        else:
            academic_year = year or "2026 - 2030"

        cur.execute(
            """
            INSERT INTO courses (
                course_code, course_name, description, credits, department_id,
                faculty_id, type, semester, academic_year, created_at
            ) VALUES (%s,%s,NULL,%s,%s,NULL,%s,%s,%s,NOW())
            ON DUPLICATE KEY UPDATE
                course_name=VALUES(course_name),
                credits=VALUES(credits),
                type=VALUES(type),
                academic_year=VALUES(academic_year)
            """,
            (
                code.upper(),
                name,
                max(0, credits),
                dep_id,
                normalize_course_type(r.get("CourseType")),
                semester,
                academic_year,
            ),
        )
        ok += 1

    return ok, skip


def import_faculty_courses(cur):
    rows = get_sheet_rows("Faculty_Courses.xlsx")
    ok, skip = 0, 0

    for r in rows:
        faculty_id = norm(r.get("FacultyID"))
        course_code = norm(r.get("CourseCode"))
        if not faculty_id or not course_code:
            skip += 1
            continue

        cur.execute("SELECT 1 FROM faculty WHERE faculty_id=%s", (faculty_id,))
        if not cur.fetchone():
            skip += 1
            continue

        cur.execute(
            """
            INSERT INTO faculty_courses (faculty_id, course_code)
            VALUES (%s, %s)
            ON DUPLICATE KEY UPDATE course_code=VALUES(course_code)
            """,
            (faculty_id, course_code.upper()),
        )
        ok += 1

    return ok, skip


def import_rooms(cur):
    rows = get_sheet_rows("Rooms.xlsx")
    ok_rooms, ok_alloc, skip = 0, 0, 0

    for r in rows:
        room_no = norm(r.get("RoomNo"))
        room_type = normalize_room_type(r.get("Type"))
        department = normalize_dept_code(r.get("Department"))
        semester = norm_int(r.get("Semester"))
        section = norm(r.get("Section"))
        course_code = norm(r.get("CourseCode"))

        if not room_no:
            skip += 1
            continue

        cur.execute(
            """
            INSERT INTO rooms (
                room_number, room_type, assigned_department, assigned_semester,
                building, capacity, is_available, created_at
            ) VALUES (%s,%s,%s,%s,NULL,NULL,1,NOW())
            ON DUPLICATE KEY UPDATE
                room_type=VALUES(room_type),
                assigned_department=VALUES(assigned_department),
                assigned_semester=VALUES(assigned_semester)
            """,
            (room_no, room_type, department, semester),
        )
        ok_rooms += 1

        cur.execute(
            """
            INSERT INTO room_allocations (
                room_number, room_type, department, semester, section, course_code
            ) VALUES (%s,%s,%s,%s,%s,%s)
            ON DUPLICATE KEY UPDATE
                room_type=VALUES(room_type)
            """,
            (room_no, room_type, department, semester, section, (course_code or "").upper() if course_code else None),
        )
        ok_alloc += 1

    return ok_rooms, ok_alloc, skip


def main():
    started = datetime.now()
    conn = pymysql.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cur:
            ensure_schema(cur)

            f_ok, f_skip = import_faculty(cur)
            s_ok, s_skip = import_students(cur)
            c_ok, c_skip = import_courses(cur)
            fc_ok, fc_skip = import_faculty_courses(cur)
            r_ok, ra_ok, r_skip = import_rooms(cur)

        conn.commit()

        print("IMPORT COMPLETE")
        print(f"Faculty imported/upserted: {f_ok}, skipped: {f_skip}")
        print(f"Students imported/upserted: {s_ok}, skipped: {s_skip}")
        print(f"Courses imported/upserted: {c_ok}, skipped: {c_skip}")
        print(f"Faculty-Courses imported/upserted: {fc_ok}, skipped: {fc_skip}")
        print(f"Rooms upserted: {r_ok}, Room allocations upserted: {ra_ok}, skipped: {r_skip}")
        print(f"Duration: {datetime.now() - started}")

    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
