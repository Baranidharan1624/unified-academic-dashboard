import re
from pathlib import Path

import bcrypt
import mysql.connector
from openpyxl import load_workbook

ROOT = Path(r"d:\projects\unified-academic-dashboard-Hari\excels")
DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "12345",
    "database": "campusone",
    "autocommit": False,
}


def normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def read_sheet(path: Path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    hmap = {normalize_header(h): i for i, h in enumerate(headers)}

    rows = []
    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rows.append(values)

    wb.close()
    return hmap, rows


def get(values, hmap, *aliases):
    for a in aliases:
        idx = hmap.get(normalize_header(a))
        if idx is None or idx >= len(values):
            continue
        v = values[idx]
        if v is None:
            continue
        s = str(v).strip()
        if s != "":
            return s
    return None


def to_int(value):
    if value is None:
        return None
    s = str(value).strip()
    if s == "":
        return None
    try:
        return int(float(s))
    except Exception:
        return None


def normalize_academic_year(value):
    if value is None:
        return None
    s = str(value).strip()
    if re.fullmatch(r"\d{4}", s):
        start = int(s)
        return f"{start}-{start + 4}"
    if re.fullmatch(r"\d{4}\s*-\s*\d{4}", s):
        a, b = [x.strip() for x in s.split("-")]
        return f"{a}-{b}"
    return s


def normalize_course_year(value):
    if value is None:
        return None
    s = str(value).strip()
    if re.fullmatch(r"\d{4}", s):
        start = int(s)
        return f"{start} - {start + 4}"
    if re.fullmatch(r"\d{4}\s*-\s*\d{4}", s):
        a, b = [x.strip() for x in s.split("-")]
        return f"{a} - {b}"
    return s


def normalize_type(value):
    if value is None:
        return "THEORY"
    s = str(value).strip().upper()
    if s in ("LAB",):
        return "LAB"
    return "THEORY"


def normalize_dept_token(value):
    if not value:
        return None
    token = str(value).strip()
    base = token.split("-", 1)[0].strip() if "-" in token else token
    upper = base.upper()
    if upper in ("BIOTECH", "BIOTECHNOLOGY"):
        return "BIO"
    return base


def ensure_schema(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS faculty_courses (
            faculty_id VARCHAR(255) NOT NULL,
            course_code VARCHAR(255) NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (faculty_id, course_code),
            INDEX idx_faculty_courses_course (course_code),
            CONSTRAINT fk_faculty_courses_faculty
              FOREIGN KEY (faculty_id) REFERENCES faculty(faculty_id)
              ON DELETE CASCADE
        ) ENGINE=InnoDB
        """
    )

    cur.execute("SHOW COLUMNS FROM students LIKE 'section'")
    if cur.fetchone() is None:
        cur.execute("ALTER TABLE students ADD COLUMN section VARCHAR(255) NULL")

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS room_course_allocations (
            id BIGINT PRIMARY KEY AUTO_INCREMENT,
            room_id BIGINT NOT NULL,
            course_code VARCHAR(255) NOT NULL,
            department VARCHAR(255),
            semester INT,
            section VARCHAR(50),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE KEY uq_room_course_allocation (room_id, course_code, department, semester, section),
            FOREIGN KEY (room_id) REFERENCES rooms(id) ON DELETE CASCADE
        ) ENGINE=InnoDB
        """
    )


def fetch_department_maps(cur):
    cur.execute("SELECT id, department_code, department_name FROM departments")
    rows = cur.fetchall()
    by_code = {}
    by_name = {}
    for did, code, name in rows:
        if code:
            by_code[str(code).strip().upper()] = did
        if name:
            by_name[str(name).strip().upper()] = did
    return by_code, by_name


def resolve_department_id(raw, by_code, by_name):
    token = normalize_dept_token(raw)
    if not token:
        return None
    up = token.upper()
    if up in by_code:
        return by_code[up]
    if up in by_name:
        return by_name[up]
    return None


def next_faculty_id(cur):
    seq = 1
    while True:
        candidate = f"FAC-{seq:04d}"
        cur.execute("SELECT 1 FROM faculty WHERE faculty_id=%s", (candidate,))
        if cur.fetchone() is None:
            return candidate
        seq += 1


def next_student_reg(cur):
    import datetime

    year = datetime.datetime.now().year
    seq = 1
    while True:
        candidate = f"STU-{year}-{seq:04d}"
        cur.execute("SELECT 1 FROM students WHERE registration_number=%s", (candidate,))
        if cur.fetchone() is None:
            return candidate
        seq += 1


def upsert_account(cur, *, full_name, email, role, department=None, semester=None, academic_year=None, age=None, mobile=None, address=None, blood_group=None):
    cur.execute("SELECT id FROM accounts WHERE email=%s", (email,))
    row = cur.fetchone()
    if row:
        account_id = row[0]
        cur.execute(
            """
            UPDATE accounts
               SET full_name=%s, role=%s, status='ACTIVE', department=%s, semester=%s, academic_year=%s,
                   age=%s, mobile_number=%s, address=%s, blood_group=%s
             WHERE id=%s
            """,
            (full_name, role, department, semester, academic_year, age, mobile, address, blood_group, account_id),
        )
        return account_id, False

    pwd_hash = bcrypt.hashpw("Campus@123".encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
    cur.execute(
        """
        INSERT INTO accounts(full_name,email,password,role,status,department,semester,academic_year,age,mobile_number,address,blood_group)
        VALUES(%s,%s,%s,%s,'ACTIVE',%s,%s,%s,%s,%s,%s,%s)
        """,
        (full_name, email, pwd_hash, role, department, semester, academic_year, age, mobile, address, blood_group),
    )
    return cur.lastrowid, True


def main():
    conn = mysql.connector.connect(**DB_CONFIG)
    cur = conn.cursor()

    try:
        cur.execute("SET SESSION innodb_lock_wait_timeout = 120")
        ensure_schema(cur)
        conn.commit()
        by_code, by_name = fetch_department_maps(cur)

        summary = {}

        # FACULTY
        hmap, rows = read_sheet(ROOT / "Faculty.xlsx")
        created = updated = skipped = 0
        for values in rows:
            email = get(values, hmap, "Email")
            full_name = get(values, hmap, "Name")
            if not email or not full_name:
                skipped += 1
                continue

            dept = normalize_dept_token(get(values, hmap, "Department"))
            faculty_id = get(values, hmap, "FacultyID")
            age = to_int(get(values, hmap, "Age"))
            mobile = get(values, hmap, "Mobile")
            address = get(values, hmap, "Address")
            blood = get(values, hmap, "BloodGroup", "Blood Group")

            account_id, is_created = upsert_account(
                cur,
                full_name=full_name,
                email=email.lower(),
                role="FACULTY",
                department=dept,
                age=age,
                mobile=mobile,
                address=address,
                blood_group=blood,
            )
            if is_created:
                created += 1
            else:
                updated += 1

            if not faculty_id:
                faculty_id = next_faculty_id(cur)

            cur.execute("SELECT account_id FROM faculty WHERE account_id=%s", (account_id,))
            if cur.fetchone():
                cur.execute("UPDATE faculty SET faculty_id=%s WHERE account_id=%s", (faculty_id, account_id))
            else:
                cur.execute("INSERT INTO faculty(account_id, faculty_id) VALUES(%s,%s)", (account_id, faculty_id))

        summary["faculty"] = {"created": created, "updated": updated, "skipped": skipped, "rows": len(rows)}
        conn.commit()

        # STUDENTS
        hmap, rows = read_sheet(ROOT / "students.xlsx")
        created = updated = skipped = 0
        for values in rows:
            email = get(values, hmap, "Email")
            full_name = get(values, hmap, "Name")
            if not email or not full_name:
                skipped += 1
                continue

            dept = normalize_dept_token(get(values, hmap, "Department"))
            semester = to_int(get(values, hmap, "Semester"))
            section = get(values, hmap, "Section")
            academic_year = normalize_academic_year(get(values, hmap, "AcademicYear", "Academic Year"))
            reg_no = get(values, hmap, "RegNo", "Reg No", "RegistrationNumber")
            age = to_int(get(values, hmap, "Age"))
            mobile = get(values, hmap, "Mobile")
            address = get(values, hmap, "Address")
            blood = get(values, hmap, "BloodGroup", "Blood Group")

            account_id, is_created = upsert_account(
                cur,
                full_name=full_name,
                email=email.lower(),
                role="STUDENT",
                department=dept,
                semester=semester,
                academic_year=academic_year,
                age=age,
                mobile=mobile,
                address=address,
                blood_group=blood,
            )
            if is_created:
                created += 1
            else:
                updated += 1

            if not reg_no:
                reg_no = next_student_reg(cur)

            if not section:
                section = "A"

            cur.execute("SELECT account_id FROM students WHERE account_id=%s", (account_id,))
            if cur.fetchone():
                cur.execute(
                    "UPDATE students SET registration_number=%s, section=%s WHERE account_id=%s",
                    (reg_no, section, account_id),
                )
            else:
                cur.execute(
                    "INSERT INTO students(account_id, registration_number, section) VALUES(%s,%s,%s)",
                    (account_id, reg_no, section),
                )

        summary["students"] = {"created": created, "updated": updated, "skipped": skipped, "rows": len(rows)}
        conn.commit()

        # FACULTY COURSES
        hmap, rows = read_sheet(ROOT / "Faculty_Courses.xlsx")
        inserted = skipped = 0
        for values in rows:
            faculty_id = get(values, hmap, "FacultyID")
            course_code = get(values, hmap, "CourseCode")
            if not faculty_id or not course_code:
                skipped += 1
                continue

            cur.execute("SELECT 1 FROM faculty WHERE faculty_id=%s", (faculty_id,))
            if cur.fetchone() is None:
                skipped += 1
                continue

            cur.execute(
                """
                INSERT INTO faculty_courses(faculty_id, course_code)
                VALUES(%s,%s)
                ON DUPLICATE KEY UPDATE course_code=VALUES(course_code)
                """,
                (faculty_id, course_code.upper()),
            )
            inserted += 1

        summary["faculty_courses"] = {"upserted": inserted, "skipped": skipped, "rows": len(rows)}
        conn.commit()

        # ROOMS + ALLOCATION
        hmap, rows = read_sheet(ROOT / "Rooms.xlsx")
        room_upserted = alloc_upserted = skipped = 0
        for values in rows:
            room_no = get(values, hmap, "RoomNo", "Room No")
            if not room_no:
                skipped += 1
                continue

            room_type = normalize_type(get(values, hmap, "Type"))
            dept = normalize_dept_token(get(values, hmap, "Department"))
            sem = to_int(get(values, hmap, "Semester"))
            section = get(values, hmap, "Section")
            course_code = get(values, hmap, "CourseCode")

            cur.execute(
                """
                INSERT INTO rooms(room_number, room_type, assigned_department, assigned_semester, is_available)
                VALUES(%s,%s,%s,%s,b'1')
                ON DUPLICATE KEY UPDATE
                    room_type=VALUES(room_type),
                    assigned_department=VALUES(assigned_department),
                    assigned_semester=VALUES(assigned_semester)
                """,
                (room_no, room_type, dept, sem),
            )
            room_upserted += 1

            cur.execute("SELECT id FROM rooms WHERE room_number=%s", (room_no,))
            room_id = cur.fetchone()[0]

            if course_code:
                cur.execute(
                    """
                    INSERT INTO room_course_allocations(room_id, course_code, department, semester, section)
                    VALUES(%s,%s,%s,%s,%s)
                    ON DUPLICATE KEY UPDATE course_code=VALUES(course_code)
                    """,
                    (room_id, course_code.upper(), dept, sem, section),
                )
                alloc_upserted += 1

        summary["rooms"] = {"room_upserted": room_upserted, "alloc_upserted": alloc_upserted, "skipped": skipped, "rows": len(rows)}
        conn.commit()

        # COURSES
        hmap, rows = read_sheet(ROOT / "course.xlsx")
        upserted = skipped = 0
        for values in rows:
            code = get(values, hmap, "CourseCode", "Course Code")
            name = get(values, hmap, "CourseName", "Course Name")
            dept_raw = get(values, hmap, "Department")
            sem = to_int(get(values, hmap, "Semester"))
            year = normalize_course_year(get(values, hmap, "Year"))
            credits = to_int(get(values, hmap, "Credits"))
            ctype = normalize_type(get(values, hmap, "CourseType", "Course Type"))

            if not code or not name or sem is None or not dept_raw or not year:
                skipped += 1
                continue
            if credits is None:
                credits = 0

            dept_id = resolve_department_id(dept_raw, by_code, by_name)
            if dept_id is None:
                skipped += 1
                continue

            cur.execute(
                """
                INSERT INTO courses(course_code, course_name, credits, type, semester, academic_year, department_id)
                VALUES(%s,%s,%s,%s,%s,%s,%s)
                ON DUPLICATE KEY UPDATE
                    course_name=VALUES(course_name),
                    credits=VALUES(credits),
                    type=VALUES(type),
                    academic_year=VALUES(academic_year)
                """,
                (code.upper(), name, credits, ctype, sem, year, dept_id),
            )
            upserted += 1

        summary["courses"] = {"upserted": upserted, "skipped": skipped, "rows": len(rows)}

        conn.commit()

        print("IMPORT_COMPLETED")
        for k, v in summary.items():
            print(k, v)

    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()


if __name__ == "__main__":
    main()
